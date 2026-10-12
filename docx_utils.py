from pathlib import Path

from docx import Document


# ============================================================
# WORD DOCUMENT TEXT EXTRACTION
# ============================================================

def extract_docx_text(file_path):
    """
    Extract text from a DOCX document.
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"DOCX file not found: {file_path}"
        )

    try:

        document = Document(
            str(file_path)
        )

        extracted = []

        # ----------------------------------------------------
        # Paragraphs
        # ----------------------------------------------------

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                extracted.append(text)

        # ----------------------------------------------------
        # Tables
        # ----------------------------------------------------

        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:
                    extracted.append(
                        " | ".join(row_text)
                    )

        return "\n".join(extracted).strip()

    except Exception as error:

        raise RuntimeError(
            f"Unable to process DOCX: {error}"
        )