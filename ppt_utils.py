from pathlib import Path

from pptx import Presentation


# ============================================================
# POWERPOINT TEXT EXTRACTION
# ============================================================

def extract_ppt_text(file_path):
    """
    Extract text from a PowerPoint presentation.

    Parameters
    ----------
    file_path : str or Path
        Path of PPTX file.

    Returns
    -------
    str
        Extracted slide text.
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"PowerPoint file not found: {file_path}"
        )

    try:

        presentation = Presentation(
            str(file_path)
        )

        extracted = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            extracted.append(
                f"\n--- Slide {slide_number} ---\n"
            )

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:
                        extracted.append(text)

        return "\n".join(extracted).strip()

    except Exception as error:

        raise RuntimeError(
            f"Unable to process PowerPoint: {error}"
        )